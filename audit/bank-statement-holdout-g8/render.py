"""Renders the G8 holdout into its frozen assets and writes the seal.

Statements are typeset by Typst or drawn by fpdf2 as each story says.
Cashbooks are CSV (UTF-8, CRLF); K02's also as a typed XLSX from XlsxWriter
(dates and amounts as numbers) and K08's as a text XLSX from openpyxl. The
committed files under frozen/ are the assets, sealed by INPUT-MANIFEST.json;
this file is their recipe. The facts are not stored: `python3 -B
stories.py` regenerates them byte for byte, and the seal carries their hash,
the producers, the font files and the eligibility declared in holdout.json.

Amiri is not a system font here. It comes from the npm package
@fontsource/amiri 5.3.0 (SIL Open Font License): its Arabic and Latin
subsets at weight 400 merged with fontTools into one TrueType file, passed
as --amiri. Its hash is recorded with the other fonts.

  PYTHONPATH=<typst, fpdf2, uharfbuzz, fonttools, xlsxwriter, openpyxl> \\
    python3 -B render.py --amiri <Amiri merged .ttf> [--out dir] [--work dir]
"""
import argparse
import hashlib
import json
import sys
import tempfile
from datetime import date, datetime
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import render_fpdf  # noqa: E402
import render_typst  # noqa: E402
from common import CASHBOOK_HEADERS, cashbook_csv, cashbook_rows  # noqa: E402
from stories import VERSION, build, facts_json  # noqa: E402

SYSTEM_FONTS = {
    'DejaVu Serif': '/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf',
    'DejaVu Serif Bold': '/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf',
    'Inter': '/usr/share/fonts/opentype/inter/Inter-Regular.otf',
    'DejaVu Sans Mono': '/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf',
    'DejaVu Sans': '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
    'FreeSans': '/usr/share/fonts/truetype/freefont/FreeSans.ttf',
    'Caladea': '/usr/share/fonts/truetype/crosextra/Caladea-Regular.ttf',
    'Carlito': '/usr/share/fonts/truetype/crosextra/Carlito-Regular.ttf',
}


def sha(b):
    return hashlib.sha256(b).hexdigest()


def typed_xlsx(story, path):
    import xlsxwriter
    book = xlsxwriter.Workbook(str(path), {'constant_memory': False})
    book.set_properties({'created': date(2026, 10, 23)})
    sheet = book.add_worksheet('Cashbook')
    day = book.add_format({'num_format': 'yyyy-mm-dd'})
    money = book.add_format({'num_format': '#,##0.00'})
    sheet.write_row(0, 0, CASHBOOK_HEADERS)
    for i, e in enumerate(story['cashbook'], 1):
        sheet.write_datetime(i, 0, date.fromisoformat(e['date']), day)
        sheet.write_string(i, 1, e['voucher'])
        sheet.write_string(i, 2, e['description'])
        if e['reference']:
            sheet.write_string(i, 3, e['reference'])
        if e['receipt']:
            sheet.write_number(i, 4, e['receipt'] / 100, money)
        if e['payment']:
            sheet.write_number(i, 5, e['payment'] / 100, money)
    book.close()


def text_xlsx(story, path):
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = 'Cashbook'
    for row in cashbook_rows(story):
        sheet.append([cell if cell else None for cell in row])
    for row in sheet.iter_rows():
        for cell in row:
            cell.number_format = '@'
    book.properties.created = book.properties.modified = datetime(2026, 10, 23)
    book.save(str(path))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--out', default=str(HERE / 'frozen'))
    parser.add_argument('--work', default=None)
    parser.add_argument('--amiri', required=True)
    args = parser.parse_args()
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    amiri = Path(args.amiri)
    stories = build()
    facts = facts_json(stories)
    written = []
    for s in stories:
        path = out / f"{s['id']}-cashbook.csv"
        path.write_bytes(cashbook_csv(s).encode('utf-8'))
        written.append(path)
        if s.get('cashbookXlsx'):
            path = out / f"{s['id']}-cashbook.xlsx"
            (typed_xlsx if s['cashbookXlsx'] == 'typed' else text_xlsx)(s, path)
            written.append(path)
    with tempfile.TemporaryDirectory() as tmp:
        work = Path(args.work or tmp)
        written += render_typst.render(stories, out, work, [amiri.parent])
    written += render_fpdf.render(stories, out, amiri)
    from importlib.metadata import version as installed
    holdout = (HERE / 'holdout.json').read_bytes()
    fonts = {name: {'file': path, 'sha256': sha(Path(path).read_bytes())} for name, path in SYSTEM_FONTS.items()}
    fonts['Amiri'] = {'file': '@fontsource/amiri 5.3.0, amiri-arabic-400-normal.woff + amiri-latin-400-normal.woff, '
                              'merged by fontTools', 'sha256': sha(amiri.read_bytes())}
    manifest = {
        'version': VERSION,
        'producers': {
            **render_typst.version(),
            **render_fpdf.version(),
            'xlsxwriter': installed('xlsxwriter'),
            'openpyxl': installed('openpyxl'),
            'fonttools': installed('fonttools'),
            'cashbookCsv': 'Python csv, UTF-8, CRLF',
        },
        'fonts': fonts,
        'holdout': {'file': 'holdout.json', 'sha256': sha(holdout)},
        'facts': {'file': 'stories.json', 'stored': False,
                  'recipe': 'python3 -B audit/bank-statement-holdout-g8/stories.py',
                  'sha256': sha(facts.encode('utf-8')), 'bytes': len(facts.encode('utf-8'))},
        'files': [{'file': p.name, 'sha256': sha(p.read_bytes()), 'bytes': p.stat().st_size}
                  for p in sorted(written, key=lambda p: p.name)],
    }
    (out / 'INPUT-MANIFEST.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=1) + '\n', encoding='utf-8')
    print(json.dumps({'files': len(manifest['files']), 'producers': manifest['producers']}, ensure_ascii=False))


if __name__ == '__main__':
    main()
