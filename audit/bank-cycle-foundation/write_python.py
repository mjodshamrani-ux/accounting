"""Write one M4 cashbook with openpyxl 3.1.5 or XlsxWriter 3.2.9.

  python3.13 -B write_python.py spec.json out.xlsx

The spec is what stories.py describes: a writer, and sheets of rows. A cell
is a string, or {"n": text} for a typed number, optionally with "fmt" for
its number format. Merges are A1 ranges; hidden columns are 1-based. Both
libraries write their own files: shared strings, styles and document
properties as each makes them. Run under python3.13, where both are
installed; the generator itself needs neither.
"""
import json
import sys
from decimal import Decimal


def number(text):
    d = Decimal(text)
    return int(d) if d == d.to_integral_value() and '.' not in text else float(d)


def with_openpyxl(spec, out):
    import openpyxl
    from openpyxl.utils import get_column_letter
    book = openpyxl.Workbook()
    book.remove(book.active)
    for s in spec['sheets']:
        ws = book.create_sheet(s['name'])
        for r, row in enumerate(s['rows'], start=1):
            for c, cell in enumerate(row, start=1):
                if isinstance(cell, dict):
                    target = ws.cell(row=r, column=c, value=number(cell['n']))
                    if cell.get('fmt'):
                        target.number_format = cell['fmt']
                elif cell != '':
                    ws.cell(row=r, column=c, value=cell)
        for m in s.get('merges', []):
            ws.merge_cells(m)
        for c in s.get('hidden', []):
            ws.column_dimensions[get_column_letter(c)].hidden = True
    book.properties.creator = 'Tarasuf synthetic M4'
    book.save(out)


def with_xlsxwriter(spec, out):
    import xlsxwriter
    book = xlsxwriter.Workbook(out)
    book.set_properties({'author': 'Tarasuf synthetic M4'})
    formats = {}
    for s in spec['sheets']:
        ws = book.add_worksheet(s['name'])
        for r, row in enumerate(s['rows']):
            for c, cell in enumerate(row):
                if isinstance(cell, dict):
                    fmt = None
                    if cell.get('fmt'):
                        fmt = formats.setdefault(cell['fmt'], book.add_format({'num_format': cell['fmt']}))
                    ws.write_number(r, c, number(cell['n']), fmt)
                elif cell != '':
                    ws.write_string(r, c, cell)
        for m in s.get('merges', []):
            # A merge holds the text of its first cell, as the rows give it.
            first = m.split(':')[0]
            col = ord(first[0]) - 65
            row = int(first[1:]) - 1
            text = s['rows'][row][col] if row < len(s['rows']) and col < len(s['rows'][row]) else ''
            ws.merge_range(m, text if isinstance(text, str) else '')
        for c in s.get('hidden', []):
            ws.set_column(c - 1, c - 1, None, None, {'hidden': True})
    book.close()


def main():
    spec = json.loads(open(sys.argv[1], encoding='utf-8').read())
    out = sys.argv[2]
    {'openpyxl': with_openpyxl, 'xlsxwriter': with_xlsxwriter}[spec['writer']](spec, out)


if __name__ == '__main__':
    main()
