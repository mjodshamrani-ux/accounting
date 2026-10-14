"""Private independent source draft. No product/oracle imports or expected outputs."""
import argparse
import base64
import csv
import hashlib
import io
import json
import re
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile, ZipInfo

SCOPE = [
    ('Entity','Synthetic INV1C Owner'), ('Inventory ledger','STOCK-MATRIX'),
    ('GL ledger','GL-MATRIX'), ('Currency','KWD'), ('Currency basis','functional'),
    ('Posting status','posted'), ('Posting layer','Current'), ('As of','2026-10-31'),
    ('Chart version','CHART-MX'), ('Map version','MAP-MX'),
    ('Valuation policy version','SUPPLIED-MX'), ('Inventory account set version','INV-MX'),
]
FIELDS = ['Item ID','Item dimensions','Quantity','Unit','Quantity precision','Posted value','Valuation reference']
ITEMS = [
    ['0014','plant=P1;batch=K1','1.25','KG','2','123.456','VAL-0014-P1-K1','INV-RM','plant=P1;branch=10'],
    ['0014','plant=P2;batch=K2','2.50','KG','2','210.111','VAL-0014-P2-K2','INV-RM','plant=P2;branch=20'],
    ['014','plant=P1;batch=E1','3','EA','0','333.567','VAL-014-P1-E1','INV-PACK','plant=P1;branch=10'],
    ['14','plant=P2;batch=Z0','0','EA','0','0.000','VAL-14-P2-Z0','INV-PACK','plant=P2;branch=20'],
]


def companions():
    headers = [
        ['Mapping ID','Item ID','Item dimensions','GL account','GL dimensions','Evidence ID'],
        ['Evidence ID','Item ID','Item dimensions','Unit','Quantity precision','Valuation reference','GL account','GL dimensions','Account class','Sign','Value basis','Valid from','Valid to','Reference'],
        ['GL account','GL dimensions','Account class','Debit','Credit'],
    ]
    sources = [[],[],[]]
    for number, item in enumerate(ITEMS):
        letter = chr(65 + number)
        code, dims, _, unit, precision, amount, reference, account, gl_dims = item
        sources[0].append([f'MAP-MX-{letter}',code,dims,account,gl_dims,f'POL-MX-{letter}'])
        sources[1].append([f'POL-MX-{letter}',code,dims,unit,precision,reference,account,gl_dims,'inventory','debit-positive','provided-posted-carrying-value','2026-01-01','2026-12-31',f'SYN-MX-{letter}'])
        sources[2].append([account,gl_dims,'inventory',amount,'0.000'])
    return [[head + [k for k,_ in SCOPE]] + [row + [v for _,v in SCOPE] for row in records] for head,records in zip(headers,sources)]


def matrix(bands=((0,1,2,3),), roles=(0,1,2,3,4,5,6), separator=False, control=False):
    rows = [['Format','INV1A-MATRIX-1']] + [['Scope',k,v] for k,v in SCOPE]
    for band in bands:
        columns = [f'C{n+1:02}' for n in range(len(band))]
        if separator: columns.insert(1,'')
        rows.append(['Field'] + columns)
        for role in roles:
            values = [ITEMS[n][role] for n in band]
            if separator: values.insert(1,'')
            rows.append([FIELDS[role]] + values)
    if control: rows.append(['Report total','667.134'])
    return rows + [['End report']]


def variants():
    yield 'INV1C-P01','csv','stdlib-csv','single-band',matrix()
    yield 'INV1C-P02','xlsx','xlsxwriter','single-band',matrix()
    yield 'INV1C-P03','csv','stdlib-csv','completed-split-bands',matrix(bands=((0,2),(1,3)))
    yield 'INV1C-P04','xlsx','openpyxl','completed-split-bands',matrix(bands=((0,2),(1,3)))
    yield 'INV1C-P05','csv','stdlib-csv','empty-band-separator-column',matrix(separator=True)
    yield 'INV1C-P06','xlsx','xlsxwriter','role-axis-permutation',matrix(bands=((3,2,1,0),),roles=(6,4,3,5,1,2,0),control=True)
    rows = matrix(); next(r for r in rows if r[0]=='Posted value')[1:3] = ['210.111','123.456']
    yield 'INV1C-N01','csv','stdlib-csv','money-swap-preserving-account-grand-totals',rows
    rows = matrix(); next(r for r in rows if r[0]=='Item dimensions')[1:3] = [ITEMS[1][1],ITEMS[0][1]]
    yield 'INV1C-N02','xlsx','xlsxwriter','dimension-swap-with-original-refs',rows
    rows = matrix(bands=((0,2),(1,3))); [r for r in rows if r[0]=='Posted value'][1][2]=''
    yield 'INV1C-N03','csv','stdlib-csv','missing-zero-value',rows
    rows = matrix(bands=((0,2),(1,3))); index=next(i for i,r in enumerate(rows) if i>22 and r[0]=='Quantity'); rows.insert(index+1,rows[index].copy())
    yield 'INV1C-N04','csv','stdlib-csv','duplicate-last-band-role',rows
    rows = matrix(bands=((0,2),(1,3))); [r for r in rows if r[0]=='Item dimensions'][1][1] = ITEMS[0][1]
    yield 'INV1C-N05','csv','stdlib-csv','duplicate-global-item-key-across-bands',rows
    rows = matrix(separator=True); next(r for r in rows if r[0]=='Posted value')[2] = '0.000'
    yield 'INV1C-N06','xlsx','openpyxl','unnamed-populated-separator-column',rows
    rows = matrix(bands=((0,2),(1,3))); rows.pop(next(i for i,r in enumerate(rows) if r[0]=='Valuation reference'))
    yield 'INV1C-N07','csv','stdlib-csv','next-axis-before-prior-band-complete',rows
    yield 'INV1C-N08','xlsx','xlsxwriter','numeric-item-identity',matrix()


def csv_bytes(rows):
    stream = io.StringIO(newline=''); csv.writer(stream,lineterminator='\n').writerows(rows)
    return stream.getvalue().encode('utf-8')


def write_native(path, rows, producer, identity):
    if producer=='openpyxl':
        import openpyxl
        if openpyxl.__version__!='3.1.5': raise ValueError('PINNED_PRODUCER_VERSION')
        from openpyxl.styles import Font
        book=openpyxl.Workbook(); sheet=book.active; sheet.title='Transposed stock'
        for column in ['A','B','C','D','E','F']: sheet.column_dimensions[column].width=30
        for r,row in enumerate(rows,1):
            for c,value in enumerate(row,1):
                if value!='':
                    numeric = c>1 and row[0] in ['Quantity','Quantity precision','Posted value','Report total']
                    cell=sheet.cell(r,c,float(value) if numeric else value); cell.font=Font(name='Arial',size=12,color='FF0A1B2E')
        book.properties.created=book.properties.modified=datetime(2026,10,31); book.save(path)
    else:
        import xlsxwriter
        if xlsxwriter.__version__!='3.2.9': raise ValueError('PINNED_PRODUCER_VERSION')
        book=xlsxwriter.Workbook(path); book.set_properties({'created':datetime(2026,10,31)}); sheet=book.add_worksheet('Transposed stock'); sheet.set_column('A:F',30)
        for r,row in enumerate(rows):
            for c,value in enumerate(row):
                if value!='':
                    numeric = c>0 and row[0] in ['Quantity','Quantity precision','Posted value','Report total']
                    if identity=='INV1C-N08' and row[0]=='Item ID' and c==1: sheet.write_number(r,c,14)
                    elif numeric: sheet.write_number(r,c,float(value))
                    else: sheet.write_string(r,c,value)
        book.close()
    with ZipFile(path) as archive: parts=[(name,archive.read(name)) for name in archive.namelist()]
    with ZipFile(path,'w') as archive:
        for name,data in parts:
            if name=='docProps/core.xml': data=re.sub(rb'(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)',rb'\g<1>2026-10-31T00:00:00Z\2',data)
            archive.writestr(ZipInfo(name,(2026,10,31,0,0,0)),data)


def produce(root, regenerate=False):
    if root.exists() and any(root.iterdir()): raise ValueError('REFUSE_REPLACE')
    root.mkdir(parents=True,exist_ok=True)
    archive_path=Path(__file__).parent/'xlsx-originals.json'
    frozen = {e['case']:e for e in json.loads(archive_path.read_text())['files']} if archive_path.exists() else {}
    artifacts=[]; index=[]
    for identity,extension,producer,variant,rows in variants():
        folder=root/'cases'/identity; folder.mkdir(parents=True); target=folder/f'register.{extension}'
        if extension=='csv': target.write_bytes(csv_bytes(rows))
        elif regenerate:
            write_native(target,rows,producer,identity)
        else:
            entry=frozen[identity]; target.write_bytes(base64.b64decode(entry['base64'],validate=True))
            if hashlib.sha256(target.read_bytes()).hexdigest()!=entry['sha256']: raise ValueError('FROZEN_NATIVE_BYTES')
        if extension=='xlsx':
            raw=target.read_bytes(); artifacts.append({'case':identity,'producer':producer,'version':{'openpyxl':'3.1.5','xlsxwriter':'3.2.9'}[producer],'bytes':len(raw),'sha256':hashlib.sha256(raw).hexdigest(),'base64':base64.b64encode(raw).decode()})
        for number,source in enumerate(companions(),1): (folder/f'source-{number}.csv').write_bytes(csv_bytes(source))
        index.append({'id':identity,'container':extension,'producer':producer,'variant':variant,'files':[f'cases/{identity}/register.{extension}']+[f'cases/{identity}/source-{n}.csv' for n in [1,2,3]]})
    (root/'corpus.json').write_text(json.dumps({'schema':'inv1c-matrix-corpus/1','synthetic':True,'financialStories':1,'cases':index},indent=2)+'\n')
    if regenerate: (root/'xlsx-originals.json').write_text(json.dumps({'schema':'inv1c-native-originals/1','synthetic':True,'files':artifacts},indent=2)+'\n')


if __name__=='__main__':
    parser=argparse.ArgumentParser(); parser.add_argument('--output-dir',type=Path,required=True); parser.add_argument('--regenerate-xlsx',action='store_true'); args=parser.parse_args(); produce(args.output_dir,args.regenerate_xlsx)
