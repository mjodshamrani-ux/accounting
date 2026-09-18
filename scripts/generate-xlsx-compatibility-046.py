"""Synthetic positive compatibility fixtures from two independent XLSX writers."""
from pathlib import Path
from datetime import datetime
import json
import openpyxl
import xlsxwriter

dest = Path(__file__).resolve().parents[1] / 'tests/fixtures/xlsx-046'
dest.mkdir(parents=True, exist_ok=True)
records = [
    [datetime(2026, 7, 15), '000123', 'Invoice', 123.45, 'SAR', 'Synthetic invoice'],
    [datetime(2026, 7, 16), '12345678901234567890', 'Credit Note', -25, 'SAR', 'Synthetic credit note'],
    [datetime(2026, 7, 17), 'INV-HIDDEN', 'Invoice', 10, 'SAR', 'Synthetic hidden row retained once'],
]
headers = ['Date', 'Document No', 'Document Type', 'Amount', 'Currency', 'Description', 'Unused helper']
arabic = ['التاريخ', 'رقم المستند', 'نوع المستند', 'المبلغ', 'العملة', 'الوصف', 'مساعد غير مستخدم']
for producer in ['openpyxl', 'xlsxwriter']:
    for language in ['en', 'ar']:
        file = dest / f'{producer}-{language}.xlsx'
        names = arabic if language == 'ar' else headers
        if producer == 'openpyxl':
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = 'Transactions'
            sheet.merge_cells('A1:G1')
            sheet['A1'] = 'Synthetic compatibility statement'
            sheet.append(['Currency', 'SAR'])
            sheet.append(names)
            for i, row in enumerate(records, 4):
                sheet.append(row)
                sheet.cell(i, 1).number_format = 'yyyy-mm-dd'
                sheet.cell(i, 4).number_format = '#,##0.00;[Red](#,##0.00)'
                sheet.cell(i, 2).number_format = '@'
                sheet.cell(i, 7).value = f'=D{i}/100'
                sheet.cell(i, 7).number_format = '0.00%'
            sheet.row_dimensions[6].hidden = True
            book.create_sheet('Read me').append(['Synthetic fixture only'])
            book.save(file)
        else:
            book = xlsxwriter.Workbook(file)
            sheet = book.add_worksheet('Transactions')
            date = book.add_format({'num_format': 'yyyy-mm-dd'})
            amount = book.add_format({'num_format': '#,##0.00;[Red](#,##0.00)'})
            helper = book.add_format({'num_format': '0.00%'})
            sheet.merge_range('A1:G1', 'Synthetic compatibility statement')
            sheet.write_row(1, 0, ['Currency', 'SAR'])
            sheet.write_row(2, 0, names)
            for i, row in enumerate(records, 3):
                sheet.write_datetime(i, 0, row[0], date)
                sheet.write_string(i, 1, row[1])
                sheet.write_string(i, 2, row[2])
                sheet.write_number(i, 3, row[3], amount)
                sheet.write_row(i, 4, row[4:])
                sheet.write_formula(i, 6, f'=D{i+1}/100', helper, row[3]/100)
            sheet.set_row(5, None, None, {'hidden': True})
            book.add_worksheet('Read me').write(0, 0, 'Synthetic fixture only')
            book.close()
(dest/'provenance.json').write_text(json.dumps({'synthetic':True,'producers':{'openpyxl':openpyxl.__version__,'xlsxwriter':xlsxwriter.__version__},'expectedMinor':[12345,-2500,1000],'expectedReferences':[r[1] for r in records],'expectedDates':[r[0].strftime('%Y-%m-%d') for r in records],'hiddenSourceRow':6},indent=2)+'\n')
